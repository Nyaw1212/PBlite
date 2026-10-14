from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Project
from .takeoff import material_total, used_materials


def export_used_materials(path: str | Path, project: Project) -> None:
    materials = used_materials(project)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Used materials"

    for cell_range, value in (("A1:E1", f"Project Template Name: {project.model or 'Panel Builder Lite'}"), ("A2:E2", f"Project Name: {project.project_name}"), ("A3:E3", f"Address: {project.address}"), ("A5:E5", "Used materials")):
        sheet.merge_cells(cell_range)
        sheet[cell_range.split(":")[0]] = value

    headers = ["Description", "Catalog Number", "Qty", "Unit Price", "Total"]
    sheet.append(headers)
    for material in materials:
        sheet.append([material.description, material.catalog_number, material.quantity, material.price, material.quantity * material.price])
    sheet.append(["", "", "", "Grand Total", material_total(materials)])

    grey = PatternFill("solid", fgColor="B7B7B7")
    dark_grey = PatternFill("solid", fgColor="999999")
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = dark_grey, Font(bold=True), Alignment(horizontal="center")
    for row in (5, 6):
        for cell in sheet[row]:
            cell.fill, cell.font, cell.alignment = grey, Font(bold=True), Alignment(horizontal="center")
    total_row = 7 + len(materials)
    for cell in sheet[total_row][3:5]:
        cell.fill, cell.font = grey, Font(bold=True)
    for row in sheet.iter_rows(min_row=7, max_row=total_row, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    for row in range(7, total_row + 1):
        sheet.cell(row, 4).number_format = '$#,##0.00'
        sheet.cell(row, 5).number_format = '$#,##0.00'
    for column, width in {"A": 55, "B": 20, "C": 12, "D": 14, "E": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A7"
    workbook.save(path)
