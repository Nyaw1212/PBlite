from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _normalise_header(value: Any) -> str:
    return "".join(str(value or "").strip().lower().split())


HEADER_ALIASES = {
    "catalognumber": "CatalogNumber",
    "catalognumber": "CatalogNumber",
    "catalog": "CatalogNumber",
    "description": "Description",
    "price": "Price",
    "inventorycategory": "InventoryCategory",
    "category": "InventoryCategory",
    "type": "Type",
    "size": "Size",
    "color": "Color",
    "supplier": "Supplier",
    "vendorcode": "VendorCode",
}


def as_number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value or "").replace("$", "").replace(",", ""))
    except ValueError:
        return 0.0


def load_mml(path: str | Path, sheet_name: str = "MML") -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return []

    columns = {
        index: HEADER_ALIASES.get(_normalise_header(header), str(header or "").strip())
        for index, header in enumerate(headers)
    }
    catalog: list[dict[str, Any]] = []
    for row in rows:
        item = {columns[index]: value for index, value in enumerate(row) if index in columns}
        item["CatalogNumber"] = str(item.get("CatalogNumber", "") or "").strip()
        if not item["CatalogNumber"]:
            continue
        item["Description"] = str(item.get("Description", "") or "").strip()
        item["Price"] = as_number(item.get("Price"))
        item["InventoryCategory"] = str(item.get("InventoryCategory", "VAN") or "VAN").strip().upper()
        for key in ("Type", "Size", "Color", "Supplier", "VendorCode"):
            item[key] = str(item.get(key, "") or "").strip()
        catalog.append(item)
    return catalog


def search_materials(catalog: list[dict[str, Any]], query: str = "", kinds: tuple[str, ...] = ()) -> list[dict[str, Any]]:
    query = query.strip().lower()
    kinds = tuple(kind.upper() for kind in kinds)
    results = []
    for item in catalog:
        haystack = " ".join(str(item.get(key, "")) for key in ("CatalogNumber", "Description", "Type", "Size", "Color")).lower()
        material_type = str(item.get("Type", "")).upper()
        if query and query not in haystack:
            continue
        if kinds and material_type not in kinds:
            continue
        results.append(item)
    return results
